from pathlib import Path

import pytest

from dam.extract import default_registry, ocr
from dam.extract.docx import DocxExtractor
from dam.extract.pdf import PdfExtractor
from pdfs import make_scan_pdf, make_text_pdf

LANGS = ["fi-FI", "en-US"]


def test_pdf_text_layer(tmp_path):
    pdf = make_text_pdf(tmp_path / "text.pdf")
    got = PdfExtractor(LANGS).extract(pdf)
    assert "Kotihoito" in got.text
    assert got.ocr_used is False
    assert got.extractor == "pdf"


@pytest.mark.ocr
@pytest.mark.skipif(ocr.engine() is None, reason="no OCR engine (Apple Vision or Tesseract)")
def test_pdf_ocr_fallback(tmp_path):
    pdf = make_scan_pdf(tmp_path / "scan.pdf")
    got = PdfExtractor(LANGS).extract(pdf)
    assert "Kotihoito" in got.text
    assert got.ocr_used is True


def test_text_file(tmp_path):
    f = tmp_path / "note.txt"
    f.write_text("Hello ääkköset", encoding="utf-8")
    got = default_registry(LANGS).extract(f)
    assert got is not None and got.text == "Hello ääkköset"
    assert got.extractor == "text" and got.ocr_used is False


def test_unsupported_returns_none(tmp_path):
    f = tmp_path / "a.zip"
    f.write_bytes(b"PK")
    assert default_registry(LANGS).extract(f) is None
    assert default_registry(LANGS).for_path(f) is None


def test_docx(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("Kotihoito lasku")
    d.add_paragraph("toinen kappale")
    f = tmp_path / "a.docx"
    d.save(str(f))
    got = DocxExtractor().extract(f)
    assert got.text == "Kotihoito lasku\ntoinen kappale"
    assert default_registry(LANGS).for_path(f).name == "docx"


MSG_FIXTURE = Path(__file__).parent / "data" / "mail.msg"


@pytest.mark.skipif(not MSG_FIXTURE.exists(), reason="no mail.msg fixture")
def test_msg():
    got = default_registry(LANGS).extract(MSG_FIXTURE)
    assert got is not None and got.extractor == "msg" and got.text
