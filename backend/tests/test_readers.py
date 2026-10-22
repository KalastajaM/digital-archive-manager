"""Text readers for spreadsheets, PowerPoint, old Word and iWork files; the re-queue."""

import io
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from sqlalchemy import select, text

from dam.db.engine import MIGRATIONS_DIR, write_tx
from dam.db.tables import files
from dam.extract import default_registry
from dam.extract.base import UnsupportedContent
from dam.extract.iwork import IworkExtractor
from dam.extract.office import PptxExtractor, XlsExtractor, XlsxExtractor
from dam.extract.textutil import TextutilExtractor
from pdfs import make_text_pdf

LANGS = ["fi-FI", "en-US"]
needs_textutil = pytest.mark.skipif(
    sys.platform != "darwin" or shutil.which("textutil") is None, reason="needs macOS textutil"
)


def make_xlsx(path):
    import openpyxl

    book = openpyxl.Workbook()
    book.active.title = "Kulut"
    book.active.append(["Kotihoito", 42, None])
    book.active.append(["Sähkölasku", "Elisa"])
    book.create_sheet("Toinen").append(["salainen"])
    book.save(path)
    return path


def test_xlsx(tmp_path):
    f = make_xlsx(tmp_path / "a.xlsx")
    got = XlsxExtractor().extract(f)
    assert got.extractor == "xlsx"
    assert got.text.splitlines() == [
        "Kulut",
        "Kotihoito\t42",
        "Sähkölasku\tElisa",
        "Toinen",
        "salainen",
    ]
    assert default_registry(LANGS).for_path(f).name == "xlsx"


def test_xlsx_from_bytes_and_bad_file(tmp_path):
    f = make_xlsx(tmp_path / "a.xlsx")
    assert "Kotihoito" in XlsxExtractor().extract(f, f.read_bytes()).text
    bad = tmp_path / "b.xlsx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(UnsupportedContent, match="not a valid Excel file"):
        XlsxExtractor().extract(bad)
    ole = tmp_path / "c.xlsx"
    ole.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"x" * 20)
    with pytest.raises(UnsupportedContent, match="password-protected"):
        XlsxExtractor().extract(ole)


def test_xls(tmp_path):
    import xlwt

    book = xlwt.Workbook()
    sheet = book.add_sheet("Vanha")
    sheet.write(0, 0, "Kotihoito")
    sheet.write(0, 1, 7)
    f = tmp_path / "a.xls"
    book.save(str(f))
    got = XlsExtractor().extract(f)
    assert got.text.splitlines() == ["Vanha", "Kotihoito\t7.0"]
    assert default_registry(LANGS).for_path(f).name == "xls"
    bad = tmp_path / "b.xls"
    bad.write_bytes(b"junk")
    with pytest.raises(UnsupportedContent):
        XlsExtractor().extract(bad)


def test_pptx_slides_and_notes(tmp_path):
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Kotihoito"
    slide.placeholders[1].text = "toinen rivi"
    slide.notes_slide.notes_text_frame.text = "puheenvuoro"
    f = tmp_path / "a.pptx"
    deck.save(str(f))
    got = PptxExtractor().extract(f)
    assert got.extractor == "pptx"
    assert all(s in got.text for s in ("Kotihoito", "toinen rivi", "puheenvuoro"))
    bad = tmp_path / "b.pptx"
    bad.write_bytes(b"nope")
    with pytest.raises(UnsupportedContent, match="not a valid PowerPoint file"):
        PptxExtractor().extract(bad)


@needs_textutil
def test_doc_via_textutil(tmp_path):
    src = tmp_path / "src.txt"
    src.write_text("Kotihoito lasku", encoding="utf-8")
    subprocess.run(
        ["textutil", "-convert", "doc", "-output", str(tmp_path / "a.doc"), str(src)], check=True
    )
    got = default_registry(LANGS).extract(tmp_path / "a.doc")
    assert got is not None and got.extractor == "textutil"
    assert "Kotihoito lasku" in got.text


@needs_textutil
def test_html_via_textutil_and_garbage_doc(tmp_path):
    f = tmp_path / "a.html"
    f.write_text("<html><body><p>Kotihoito</p></body></html>", encoding="utf-8")
    assert "Kotihoito" in TextutilExtractor().extract(f).text
    bad = tmp_path / "b.doc"
    bad.write_bytes(b"\x00\x01\x02 not a document")
    try:
        TextutilExtractor().extract(bad)
    except UnsupportedContent as exc:
        assert "cannot read document" in str(exc)


def make_iwork(path, members):
    with zipfile.ZipFile(path, "w") as zf:
        for name, payload in members.items():
            zf.writestr(name, payload)
    return path


def test_pages_preview_pdf(tmp_path):
    pdf = make_text_pdf(tmp_path / "p.pdf")
    f = make_iwork(
        tmp_path / "a.pages",
        {"Index/Document.iwa": b"\x00protobuf", "QuickLook/Preview.pdf": pdf.read_bytes()},
    )
    got = IworkExtractor(LANGS).extract(f)
    assert "Kotihoito" in got.text and got.extractor == "iwork" and not got.ocr_used
    for name in ("a.numbers", "a.key"):
        assert default_registry(LANGS).for_path(tmp_path / name).name == "iwork"


def test_iwork_without_preview_or_not_a_zip(tmp_path):
    f = make_iwork(tmp_path / "a.numbers", {"Index/Document.iwa": b"x"})
    with pytest.raises(UnsupportedContent, match="without a preview"):
        IworkExtractor(LANGS).extract(f)
    bad = tmp_path / "b.key"
    bad.write_bytes(b"not a zip")
    with pytest.raises(UnsupportedContent, match="not a valid iWork file"):
        IworkExtractor(LANGS).extract(bad)


@pytest.mark.ocr
def test_keynote_preview_image_is_ocrd(tmp_path):
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (900, 200), "white")
    ImageDraw.Draw(img).text((20, 60), "Kotihoito", fill="black", font=ImageFont.load_default(60))
    buf = io.BytesIO()
    img.save(buf, "JPEG")
    f = make_iwork(tmp_path / "a.key", {"QuickLook/Thumbnail.jpg": buf.getvalue()})
    got = IworkExtractor(LANGS).extract(f)
    assert got.ocr_used and "Kotihoito" in got.text


def test_indexing_a_spreadsheet_end_to_end(app, fake_archive, tmp_path):
    src = make_xlsx(tmp_path / "x.xlsx")
    fake_archive.add("10. Financial/Banks/budget.xlsx", content=src.read_bytes(), drive_id="x1")
    app.scan()
    app.index()
    with app.engine.connect() as conn:
        state = conn.execute(select(files.c.index_state).where(files.c.drive_id == "x1")).scalar()
    assert state == "indexed"


def test_migration_requeues_unsupported_files_with_a_reader(tmp_path, new_engine):
    engine = new_engine(tmp_path / "dam.sqlite")
    config = Config()
    config.set_main_option("script_location", str(MIGRATIONS_DIR))
    with write_tx(engine) as conn:
        config.attributes["connection"] = conn
        command.upgrade(config, "0022")
        for i, (ext, state) in enumerate(
            [("xlsx", "unsupported"), ("PAGES", "unsupported"), ("zip", "unsupported"),
             ("xls", "error"), ("doc", "indexed")]
        ):  # fmt: skip
            conn.execute(
                text(
                    "INSERT INTO files (drive_id, path, path_nfc, ext, index_state, attempts, "
                    "last_error) VALUES (:d, :p, :p, :e, :s, 3, 'no text extractor')"
                ),
                {"d": f"d{i}", "p": f"f{i}", "e": ext, "s": state},
            )
        command.upgrade(config, "head")
    with engine.connect() as conn:
        rows = {
            r.drive_id: (r.index_state, r.attempts, r.last_error)
            for r in conn.execute(select(files.c.drive_id, files.c.index_state,
                                         files.c.attempts, files.c.last_error))
        }  # fmt: skip
    assert rows["d0"] == ("pending", 0, None)
    assert rows["d1"] == ("pending", 0, None)
    assert rows["d2"] == ("unsupported", 3, "no text extractor")  # zip: still no reader
    assert rows["d3"][0] == "error" and rows["d4"][0] == "indexed"  # only unsupported rows move


def test_textutil_reader_is_registered_on_macos_only():
    reader = default_registry(LANGS).for_path(Path("a.doc"))
    assert (reader is not None and reader.name == "textutil") == (sys.platform == "darwin")
