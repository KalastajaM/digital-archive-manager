"""Spreadsheets (.xlsx .xlsm .xls) and PowerPoint (.pptx) text."""

import io
import struct
import zipfile
from pathlib import Path

from dam.extract import zip_guard
from dam.extract.base import ExtractedText, UnsupportedContent

MAX_CHARS = 2_000_000  # text kept per file; the rest of a huge sheet is dropped
NOT_XLSX = "not a valid Excel file"
NOT_XLS = "not a valid old Excel file"
NOT_PPTX = "not a valid PowerPoint file"
PASSWORD_OFFICE = "password-protected Office file"
# an encrypted .xlsx/.pptx is an OLE compound file (not a zip) holding the encrypted package
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _head(path: Path, data: bytes | None) -> bytes:
    if data is not None:
        return data[:8]
    try:
        with open(path, "rb") as fh:
            return fh.read(8)
    except OSError:
        return b""


def _join(parts: list[str]) -> str:
    return "\n".join(parts)[:MAX_CHARS]


def _cell_text(value: object) -> str:
    return "" if value is None else str(value).strip()


class XlsxExtractor:
    name = "xlsx"
    suffixes = {".xlsx", ".xlsm"}

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() in self.suffixes

    def extract(self, path: Path, data: bytes | None = None) -> ExtractedText:
        import openpyxl

        zip_guard.check(path, data)
        try:
            book = openpyxl.load_workbook(
                io.BytesIO(data) if data is not None else str(path),
                read_only=True,
                data_only=True,
            )
        except (zipfile.BadZipFile, KeyError, ValueError) as exc:
            if _head(path, data) == _OLE_MAGIC:
                raise UnsupportedContent(PASSWORD_OFFICE) from exc
            raise UnsupportedContent(NOT_XLSX) from exc
        parts: list[str] = []
        size = 0
        try:
            for sheet in book.worksheets:
                parts.append(sheet.title)
                for row in sheet.iter_rows(values_only=True):
                    line = "\t".join(c for c in map(_cell_text, row) if c)
                    if line:
                        parts.append(line)
                        size += len(line)
                    if size > MAX_CHARS:
                        break
                if size > MAX_CHARS:
                    break
        finally:
            book.close()
        return ExtractedText(text=_join(parts), extractor=self.name)


class XlsExtractor:
    name = "xls"

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".xls"

    def extract(self, path: Path, data: bytes | None = None) -> ExtractedText:
        import xlrd

        try:
            book = xlrd.open_workbook(
                file_contents=data if data is not None else path.read_bytes(),
                on_demand=True,
            )
        except (xlrd.XLRDError, IndexError, struct.error) as exc:
            if "encrypted" in str(exc).lower():
                raise UnsupportedContent(PASSWORD_OFFICE) from exc
            raise UnsupportedContent(NOT_XLS) from exc
        parts: list[str] = []
        size = 0
        try:
            for sheet in book.sheets():
                parts.append(sheet.name)
                for r in range(sheet.nrows):
                    line = "\t".join(c for c in map(_cell_text, sheet.row_values(r)) if c)
                    if line:
                        parts.append(line)
                        size += len(line)
                    if size > MAX_CHARS:
                        break
                if size > MAX_CHARS:
                    break
        finally:
            book.release_resources()
        return ExtractedText(text=_join(parts), extractor=self.name)


class PptxExtractor:
    name = "pptx"

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".pptx"

    @staticmethod
    def _shape_text(shape) -> list[str]:
        out: list[str] = []
        if shape.shape_type == 6:  # a group of shapes
            for inner in shape.shapes:
                out += PptxExtractor._shape_text(inner)
        if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
            out.append(shape.text_frame.text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    out.append("\t".join(cells))
        return out

    def extract(self, path: Path, data: bytes | None = None) -> ExtractedText:
        from pptx import Presentation
        from pptx.exc import PackageNotFoundError

        zip_guard.check(path, data)
        try:
            deck = Presentation(io.BytesIO(data) if data is not None else str(path))
        except (zipfile.BadZipFile, PackageNotFoundError) as exc:
            if _head(path, data) == _OLE_MAGIC:
                raise UnsupportedContent(PASSWORD_OFFICE) from exc
            raise UnsupportedContent(NOT_PPTX) from exc
        except (KeyError, ValueError) as exc:
            raise UnsupportedContent(NOT_PPTX) from exc
        parts: list[str] = []
        for slide in deck.slides:
            for shape in slide.shapes:
                parts += self._shape_text(shape)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None and notes.text.strip():
                    parts.append(notes.text)
        return ExtractedText(text=_join(parts), extractor=self.name)
